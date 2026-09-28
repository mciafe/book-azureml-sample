from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.dml.color import RGBColor
from pathlib import Path


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE = BASE_DIR / "slide_template.pptx"
OUTFILE = BASE_DIR / "ver65_ver85_capability_comparison.pptx"


def add_box(slide, x, y, w, h, text, *, fill, line, font_size=14, bold=False, color=None, radius=True, align=PP_ALIGN.CENTER):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color or RGBColor(42, 52, 64)
    return shape


def add_circle(slide, x, y, w, h, text, *, fill, line, font_size=20):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    paragraph = tf.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    return shape


def add_arrow(slide, x, y, w, h, text="", *, fill=None, font_size=16):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or RGBColor(242, 141, 43)
    shape.line.color.rgb = fill or RGBColor(242, 141, 43)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    paragraph = tf.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    return shape


def add_line(slide, x1, y1, x2, y2, color, width=2.0):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def get_blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            return layout
    raise ValueError("Blank slide layout not found in template")


def build_slide(template_path=TEMPLATE, output_path=OUTFILE):
    prs = Presentation(str(template_path))
    if prs.slides:
        raise ValueError("Template must not contain slides")
    slide = prs.slides.add_slide(get_blank_layout(prs))

    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(247, 249, 252)

    blue = RGBColor(47, 109, 181)
    blue_lt = RGBColor(223, 236, 250)
    green = RGBColor(58, 155, 97)
    green_lt = RGBColor(224, 244, 231)
    orange = RGBColor(242, 141, 43)
    orange_lt = RGBColor(253, 235, 214)
    gray = RGBColor(128, 138, 150)
    gray_lt = RGBColor(235, 239, 243)
    dark = RGBColor(42, 52, 64)
    white = RGBColor(255, 255, 255)

    add_box(slide, Inches(0.35), Inches(0.2), Inches(12.63), Inches(0.55),
            "実装追加により、回答できる問の範囲が拡大", fill=white, line=white, font_size=24, bold=True, color=dark, radius=False)
    add_box(slide, Inches(0.6), Inches(0.85), Inches(1.4), Inches(0.45), "ver65", fill=blue, line=blue, font_size=22, bold=True, color=white)
    add_arrow(slide, Inches(5.72), Inches(0.86), Inches(1.9), Inches(0.42), "追加実装", fill=orange, font_size=18)
    add_box(slide, Inches(11.0), Inches(0.85), Inches(1.4), Inches(0.45), "ver85", fill=green, line=green, font_size=22, bold=True, color=white)

    left_x, panel_y, panel_w, panel_h = Inches(0.35), Inches(1.35), Inches(6.15), Inches(5.35)
    right_x = Inches(6.83)
    for x, line, fill in [(left_x, blue, blue_lt), (right_x, green, green_lt)]:
        add_box(slide, x, panel_y, panel_w, panel_h, "", fill=fill, line=line)

    add_box(slide, Inches(0.6), Inches(1.55), Inches(2.15), Inches(0.34), "回答できる問の範囲", fill=white, line=white, font_size=16, bold=True, color=dark, radius=False, align=PP_ALIGN.LEFT)
    add_box(slide, Inches(7.08), Inches(1.55), Inches(2.15), Inches(0.34), "回答できる問の範囲", fill=white, line=white, font_size=16, bold=True, color=dark, radius=False, align=PP_ALIGN.LEFT)

    add_box(slide, Inches(0.62), Inches(1.95), Inches(2.75), Inches(2.15), "ユーザの問集合", fill=gray_lt, line=gray, font_size=14, bold=True, color=gray)
    add_box(slide, Inches(7.1), Inches(1.95), Inches(2.75), Inches(2.15), "ユーザの問集合", fill=gray_lt, line=gray, font_size=14, bold=True, color=gray)
    add_circle(slide, Inches(1.15), Inches(2.18), Inches(1.7), Inches(1.7), "ver65\n回答可能", fill=blue, line=blue)
    add_circle(slide, Inches(7.38), Inches(2.08), Inches(2.15), Inches(2.0), "ver85\n回答可能", fill=green, line=green, font_size=21)

    add_box(slide, Inches(9.6), Inches(2.35), Inches(1.95), Inches(1.1), "追加で\n回答可能", fill=orange, line=orange, font_size=18, bold=True, color=white)
    add_line(slide, Inches(9.6), Inches(2.9), Inches(9.2), Inches(2.9), orange, width=3)
    add_line(slide, Inches(9.2), Inches(2.9), Inches(9.0), Inches(3.1), orange, width=3)
    add_line(slide, Inches(9.2), Inches(2.9), Inches(9.0), Inches(2.7), orange, width=3)

    for x, y, w, h, text, fill, line in [
        (0.72, 2.12, 0.95, 0.34, "一問", blue_lt, blue),
        (0.62, 3.75, 1.4, 0.4, "索引あり", blue_lt, blue),
        (2.28, 2.02, 1.12, 0.42, "キーワード\n一致", blue_lt, blue),
        (2.18, 3.66, 1.72, 0.48, "索引なし→不明", blue_lt, blue),
        (0.95, 4.18, 2.2, 0.44, "分解できる複数問い", blue_lt, blue),
    ]:
        add_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), text, fill=fill, line=line, font_size=11)

    for x, y, w, h, text, fill, line in [
        (7.16, 2.02, 0.95, 0.34, "一問", green_lt, green),
        (7.05, 3.75, 1.4, 0.4, "索引あり", green_lt, green),
        (8.66, 2.0, 1.35, 0.48, "類義語理解", orange_lt, orange),
        (8.72, 3.6, 1.62, 0.52, "補足情報付与", orange_lt, orange),
        (7.42, 4.18, 2.25, 0.44, "ver65の範囲も継承", green_lt, green),
    ]:
        add_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), text, fill=fill, line=line, font_size=11)

    for x, y, w, h, text in [
        (3.55, 2.18, 1.9, 0.42, "類義語理解が必要"),
        (3.45, 2.74, 2.15, 0.48, "補足情報付与が必要"),
        (3.6, 3.38, 2.2, 0.48, "分類・分解が困難"),
    ]:
        add_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), text, fill=gray_lt, line=gray, font_size=10, color=gray)

    for x, y, w, h, text in [
        (10.18, 1.98, 1.95, 0.42, "ユースケース実現"),
        (10.12, 2.55, 1.75, 0.42, "不具合"),
        (9.94, 3.12, 2.2, 0.42, "入力可能性が低い問"),
        (10.02, 3.68, 2.1, 0.48, "事前知識が必要な\n複数問い"),
    ]:
        add_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), text, fill=gray_lt, line=gray, font_size=10, color=gray)

    add_box(slide, Inches(0.6), Inches(4.7), Inches(1.8), Inches(0.32), "実装", fill=white, line=white, font_size=16, bold=True, color=dark, radius=False, align=PP_ALIGN.LEFT)
    add_box(slide, Inches(7.08), Inches(4.7), Inches(2.25), Inches(0.32), "追加した実装", fill=white, line=white, font_size=16, bold=True, color=dark, radius=False, align=PP_ALIGN.LEFT)

    flow_y = Inches(5.08)
    start_x = Inches(0.72)
    box_w = Inches(1.2)
    box_h = Inches(0.58)
    gap = Inches(0.14)
    steps = ["前処理", "RAG検索", "回答生成", "FAQ出力"]
    for idx, step in enumerate(steps):
        x = start_x + idx * (box_w + gap)
        add_box(slide, x, flow_y, box_w, box_h, step, fill=blue, line=blue, font_size=15, bold=True, color=white)
        if idx < len(steps) - 1:
            add_arrow(slide, x + box_w + Inches(0.02), flow_y + Inches(0.08), Inches(0.1), Inches(0.42), fill=blue)

    flow2_y = Inches(5.45)
    start2_x = Inches(7.18)
    for idx, step in enumerate(steps):
        x = start2_x + idx * (box_w + gap)
        add_box(slide, x, flow2_y, box_w, box_h, step, fill=green, line=green, font_size=15, bold=True, color=white)
        if idx < len(steps) - 1:
            add_arrow(slide, x + box_w + Inches(0.02), flow2_y + Inches(0.08), Inches(0.1), Inches(0.42), fill=green)

    add_box(slide, Inches(7.18), Inches(4.95), Inches(1.15), Inches(0.42), "正規化", fill=orange, line=orange, font_size=14, bold=True, color=white)
    add_box(slide, Inches(8.48), Inches(4.95), Inches(1.45), Inches(0.42), "会話変数定義", fill=orange, line=orange, font_size=13, bold=True, color=white)
    add_box(slide, Inches(10.08), Inches(4.95), Inches(1.45), Inches(0.42), "レコメンド生成", fill=orange, line=orange, font_size=13, bold=True, color=white)
    add_line(slide, Inches(7.75), Inches(5.37), Inches(7.78), Inches(5.43), orange, width=2.2)
    add_line(slide, Inches(9.2), Inches(5.37), Inches(8.6), Inches(5.43), orange, width=2.2)
    add_line(slide, Inches(10.8), Inches(5.37), Inches(10.55), Inches(5.43), orange, width=2.2)

    add_box(slide, Inches(1.0), Inches(6.65), Inches(11.3), Inches(0.45),
            "未対応の領域：ユースケース実現 / 不具合 / 想定外入力 / 分類・分解が難しい複数問い",
            fill=gray_lt, line=gray, font_size=15, bold=True, color=gray)

    prs.save(str(output_path))


if __name__ == "__main__":
    build_slide()
